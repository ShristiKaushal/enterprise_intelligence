"""
Concrete ingestion handlers: PDF, DOCX, TXT, CSV, Excel.
"""
from __future__ import annotations

import io
import json
from pathlib import Path
from typing import Optional

import pandas as pd

from app.ingestion.base import BaseIngestionHandler, IngestionResult, PageContent
from app.core.logging import get_logger

logger = get_logger(__name__)


# ── PDF Handler ───────────────────────────────────────────────────────────────

class PDFHandler(BaseIngestionHandler):
    @property
    def supported_extensions(self) -> list[str]:
        return ["pdf"]

    async def ingest(self, file_path: Path) -> IngestionResult:
        import asyncio
        import functools
        loop = asyncio.get_event_loop()
        return await loop.run_in_executor(None, functools.partial(self._ingest_sync, file_path))

    def _ingest_sync(self, file_path: Path) -> IngestionResult:
        try:
            import pdfplumber
        except ImportError:
            return IngestionResult(
                file_type="pdf", filename=file_path.name,
                total_pages=0, total_words=0, pages=[], full_text="",
                metadata={}, error="pdfplumber not installed",
            )

        pages = []
        tables_found = []
        full_text_parts = []
        doc_meta = {}

        try:
            with pdfplumber.open(file_path) as pdf:
                doc_meta = pdf.metadata or {}
                for i, page in enumerate(pdf.pages, start=1):
                    text = page.extract_text() or ""
                    raw_tables = page.extract_tables() or []
                    tables_str = []
                    for t in raw_tables:
                        tables_str.append([[str(cell) if cell else "" for cell in row] for row in t])
                        tables_found.append({
                            "page": i,
                            "rows": len(t),
                            "cols": len(t[0]) if t else 0,
                            "data": tables_str[-1],
                        })
                    pages.append(PageContent(page_number=i, text=text, tables=tables_str))
                    full_text_parts.append(text)

            full_text = "\n\n".join(full_text_parts)
            return IngestionResult(
                file_type="pdf",
                filename=file_path.name,
                total_pages=len(pages),
                total_words=self._count_words(full_text),
                pages=pages,
                full_text=full_text,
                metadata=doc_meta,
                tables=tables_found,
            )
        except Exception as e:
            logger.error("PDF ingestion failed", path=str(file_path), error=str(e))
            return IngestionResult(
                file_type="pdf", filename=file_path.name,
                total_pages=0, total_words=0, pages=[], full_text="",
                metadata={}, error=str(e),
            )


# ── DOCX Handler ──────────────────────────────────────────────────────────────

class DOCXHandler(BaseIngestionHandler):
    @property
    def supported_extensions(self) -> list[str]:
        return ["docx", "doc"]

    async def ingest(self, file_path: Path) -> IngestionResult:
        import asyncio
        import functools
        loop = asyncio.get_event_loop()
        return await loop.run_in_executor(None, functools.partial(self._ingest_sync, file_path))

    def _ingest_sync(self, file_path: Path) -> IngestionResult:
        try:
            from docx import Document
        except ImportError:
            return IngestionResult(
                file_type="docx", filename=file_path.name,
                total_pages=0, total_words=0, pages=[], full_text="",
                metadata={}, error="python-docx not installed",
            )

        try:
            doc = Document(file_path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            full_text = "\n".join(paragraphs)

            tables_found = []
            for i, table in enumerate(doc.tables):
                rows = [[cell.text for cell in row.cells] for row in table.rows]
                tables_found.append({"table_index": i, "rows": len(rows), "data": rows})

            # DOCX doesn't have strict pages — treat as one page
            pages = [PageContent(page_number=1, text=full_text, tables=[])]

            meta = {}
            if doc.core_properties:
                cp = doc.core_properties
                meta = {
                    "author": cp.author,
                    "title": cp.title,
                    "subject": cp.subject,
                    "created": str(cp.created) if cp.created else None,
                    "modified": str(cp.modified) if cp.modified else None,
                }

            return IngestionResult(
                file_type="docx",
                filename=file_path.name,
                total_pages=1,
                total_words=self._count_words(full_text),
                pages=pages,
                full_text=full_text,
                metadata=meta,
                tables=tables_found,
            )
        except Exception as e:
            logger.error("DOCX ingestion failed", path=str(file_path), error=str(e))
            return IngestionResult(
                file_type="docx", filename=file_path.name,
                total_pages=0, total_words=0, pages=[], full_text="",
                metadata={}, error=str(e),
            )


# ── TXT Handler ───────────────────────────────────────────────────────────────

class TXTHandler(BaseIngestionHandler):
    @property
    def supported_extensions(self) -> list[str]:
        return ["txt", "log", "md"]

    async def ingest(self, file_path: Path) -> IngestionResult:
        try:
            text = file_path.read_text(encoding="utf-8", errors="replace")
            lines = text.split("\n")
            # Split into virtual pages of ~100 lines each
            page_size = 100
            pages = []
            for i in range(0, len(lines), page_size):
                chunk = "\n".join(lines[i:i + page_size])
                pages.append(PageContent(page_number=i // page_size + 1, text=chunk))

            return IngestionResult(
                file_type="txt",
                filename=file_path.name,
                total_pages=len(pages),
                total_words=self._count_words(text),
                pages=pages,
                full_text=text,
                metadata={"line_count": len(lines)},
            )
        except Exception as e:
            logger.error("TXT ingestion failed", path=str(file_path), error=str(e))
            return IngestionResult(
                file_type="txt", filename=file_path.name,
                total_pages=0, total_words=0, pages=[], full_text="",
                metadata={}, error=str(e),
            )


# ── CSV Handler ───────────────────────────────────────────────────────────────

class CSVHandler(BaseIngestionHandler):
    @property
    def supported_extensions(self) -> list[str]:
        return ["csv"]

    async def ingest(self, file_path: Path) -> IngestionResult:
        import asyncio
        import functools
        loop = asyncio.get_event_loop()
        return await loop.run_in_executor(None, functools.partial(self._ingest_sync, file_path))

    def _ingest_sync(self, file_path: Path) -> IngestionResult:
        try:
            df = pd.read_csv(file_path, dtype=str, keep_default_na=True)
            return self._process_dataframe(df, file_path, "csv")
        except Exception as e:
            logger.error("CSV ingestion failed", path=str(file_path), error=str(e))
            return IngestionResult(
                file_type="csv", filename=file_path.name,
                total_pages=0, total_words=0, pages=[], full_text="",
                metadata={}, error=str(e),
            )

    def _process_dataframe(
        self, df: pd.DataFrame, file_path: Path, file_type: str
    ) -> IngestionResult:
        # Column statistics
        col_stats = {}
        for col in df.columns:
            series = df[col]
            col_stats[col] = {
                "dtype": str(series.dtype),
                "non_null_count": int(series.notna().sum()),
                "null_count": int(series.isna().sum()),
                "null_pct": round(series.isna().mean() * 100, 2),
                "unique_count": int(series.nunique()),
                "sample_values": series.dropna().head(3).tolist(),
            }

        # Raw rows as list of dicts
        raw_rows = df.where(pd.notna(df), None).to_dict(orient="records")

        # Text representation for NLP
        text_parts = []
        for _, row in df.head(50).iterrows():  # First 50 rows for text context
            text_parts.append(" | ".join(f"{k}: {v}" for k, v in row.items() if v))
        full_text = "\n".join(text_parts)

        # Virtual pages
        page_size = 50
        pages = []
        for i in range(0, len(text_parts), page_size):
            chunk = "\n".join(text_parts[i:i + page_size])
            pages.append(PageContent(page_number=i // page_size + 1, text=chunk))

        meta = {
            "row_count": len(df),
            "column_count": len(df.columns),
            "columns": list(df.columns),
            "duplicate_rows": int(df.duplicated().sum()),
        }

        return IngestionResult(
            file_type=file_type,
            filename=file_path.name,
            total_pages=len(pages),
            total_words=self._count_words(full_text),
            pages=pages if pages else [PageContent(page_number=1, text=full_text)],
            full_text=full_text,
            metadata=meta,
            raw_rows=raw_rows,
            column_stats=col_stats,
        )


class ExcelHandler(CSVHandler):
    @property
    def supported_extensions(self) -> list[str]:
        return ["xlsx", "xls"]

    def _ingest_sync(self, file_path: Path) -> IngestionResult:
        try:
            df = pd.read_excel(file_path, dtype=str)
            return self._process_dataframe(df, file_path, "xlsx")
        except Exception as e:
            logger.error("Excel ingestion failed", path=str(file_path), error=str(e))
            return IngestionResult(
                file_type="xlsx", filename=file_path.name,
                total_pages=0, total_words=0, pages=[], full_text="",
                metadata={}, error=str(e),
            )


# ── Registry ──────────────────────────────────────────────────────────────────

class IngestionHandlerRegistry:
    """Maps file extensions to their handlers."""

    def __init__(self) -> None:
        self._handlers: dict[str, BaseIngestionHandler] = {}
        # Register defaults
        for handler in [PDFHandler(), DOCXHandler(), TXTHandler(), CSVHandler(), ExcelHandler()]:
            for ext in handler.supported_extensions:
                self._handlers[ext.lower()] = handler

    def get_handler(self, extension: str) -> Optional[BaseIngestionHandler]:
        return self._handlers.get(extension.lower())

    def register(self, handler: BaseIngestionHandler) -> None:
        for ext in handler.supported_extensions:
            self._handlers[ext.lower()] = handler

    @property
    def supported_extensions(self) -> list[str]:
        return list(self._handlers.keys())


# Singleton registry
ingestion_registry = IngestionHandlerRegistry()
